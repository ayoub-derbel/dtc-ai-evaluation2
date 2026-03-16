#!/usr/bin/env python3
"""Generate AI benchmark Excel template with merged cells and formatting."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side

OUTPUT_FILENAME = "AI_Test_Benchmark.xlsx"
SHEET_NAME = "AI_Test_Benchmark"

DTCS = [
    "P0100",
    "P0110",
    "P0128",
    "P0171",
    "P0172",
    "P0300",
    "P0301",
    "P0420",
    "P0442",
    "P0455",
]

PROMPTS = [
    "prompt_v1_basic",
    "Expert_Diagnostic",
    "prompt_v2_fewshot",
    "prompt_v3_json",
]

VEHICLES = [
    "Ford Focus 2018",
    "Toyota Corolla 2018",
    "Volkswagen Golf 2019",
]


def apply_style(sheet) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    max_row = sheet.max_row
    max_col = sheet.max_column

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = sheet.cell(row=row, column=col)
            cell.border = border
            if col == 3:  # Vehicle column
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    output_path = project_root / "results" / OUTPUT_FILENAME

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    headers = ["DTC", "Prompt", "Vehicle", "GPT-4O", "Llama-3.3-70B-Versatile"]
    ws.append(headers)

    current_row = 2

    for dtc in DTCS:
        dtc_start = current_row

        for prompt in PROMPTS:
            prompt_start = current_row

            for vehicle in VEHICLES:
                ws.cell(row=current_row, column=1, value=dtc)
                ws.cell(row=current_row, column=2, value=prompt)
                ws.cell(row=current_row, column=3, value=vehicle)
                ws.cell(row=current_row, column=4, value="[Résultat à insérer]")
                ws.cell(row=current_row, column=5, value="[Résultat à insérer]")
                current_row += 1

            prompt_end = current_row - 1
            ws.merge_cells(start_row=prompt_start, start_column=2, end_row=prompt_end, end_column=2)
            ws.merge_cells(start_row=prompt_start, start_column=4, end_row=prompt_end, end_column=4)
            ws.merge_cells(start_row=prompt_start, start_column=5, end_row=prompt_end, end_column=5)

        dtc_end = current_row - 1
        ws.merge_cells(start_row=dtc_start, start_column=1, end_row=dtc_end, end_column=1)

    # Column widths
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 25
    ws.column_dimensions["E"].width = 25

    apply_style(ws)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Saved Excel file to {output_path}")


if __name__ == "__main__":
    main()
